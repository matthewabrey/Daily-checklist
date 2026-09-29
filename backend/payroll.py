"""Payroll Summary Report import — what people actually cost.

The Go2Clock file says how many hours someone worked. This one says what
those hours cost the business. Together they give a real £ per hour per
person, which is the only honest way to put money on a job record.

Matthew's definition of the cost, and the one used here:

    cost = Total Gross Pay + Employer NICs + Employer Net Pension

Employee deductions (income tax, employee NICs, student loan, accommodation,
bottled gas, training) are NOT added — they come out of the gross, they are
not on top of it. Adding them would double-count.

Two things worth remembering:

* The report is a formatted sheet, not a table. Headings sit on row 6, the
  period on row 4, and the last row is a totals row with no name. So it is
  parsed by finding the heading row and reading columns BY NAME — the same
  defence used on the work plan, which broke twice when a layout moved.
* **The payroll ID matches Go2Clock, not the app.** Checked on the
  17 Aug – 13 Sep run: 70 of 82 payroll IDs appear in Go2Clock, and eleven
  people are in the app's Name List under an older number. So people are
  matched payroll ID → Go2Clock → app, never the other way round.
"""

import logging
import re
from datetime import date

import openpyxl

logger = logging.getLogger(__name__)

# what to add up, by heading. Missing headings are treated as zero, so a
# payroll run that has no pension column still imports.
COST_PARTS = {
    "gross": ["Total Gross Pay"],
    "employer_nic": ["Employer NICs", "Employer NIC"],
    "employer_pension": ["Employer Net Pension AE", "Employer Net Pension non AE"],
}
# carried for information, never added to the cost
INFO_PARTS = {
    "net_pay": ["Net Pay"],
    "income_tax": ["Income Tax"],
    "employee_nic": ["Employee NICs", "Employee NIC"],
    "employee_pension": ["Net Pension AE", "Net Pension non AE"],
    "employer_pension_gross": ["Employer Gross Pension AE", "Employer Gross Pension non AE"],
}

_MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun",
     "jul", "aug", "sep", "oct", "nov", "dec"])}


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _num(v):
    if v is None:
        return 0.0
    try:
        return float(str(v).strip().replace(",", "").replace("£", ""))
    except ValueError:
        return 0.0


def parse_period(text):
    """'17th Aug - 13th Sep 2026' -> ('2026-08-17', '2026-09-13').

    The year is written once, at the end. A run that crosses new year
    (20th Dec - 16th Jan 2027) puts the start in the previous year.
    """
    s = str(text or "")
    yr = re.search(r"(20\d{2})", s)
    if not yr:
        return None, None
    year = int(yr.group(1))
    found = re.findall(r"(\d{1,2})\s*(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})", s)
    if len(found) < 2:
        return None, None

    def mk(d, mon, y):
        m = _MONTHS.get(mon[:3].lower())
        if not m:
            return None
        try:
            return date(y, m, int(d)).isoformat()
        except ValueError:
            return None

    end = mk(found[-1][0], found[-1][1], year)
    start = mk(found[0][0], found[0][1], year)
    if start and end and start > end:          # crossed the new year
        start = mk(found[0][0], found[0][1], year - 1)
    return start, end


def parse_payroll_summary(content: bytes):
    """-> (people, period_start, period_end, warnings)."""
    wb = openpyxl.load_workbook(__import__("io").BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    warnings = []

    # --- the heading row: the one carrying both an ID and a name column ---
    head_i, head = None, None
    for i, r in enumerate(rows[:30]):
        cells = [_norm(c) for c in (r or [])]
        if "payrollid" in cells and "employee" in cells:
            head_i, head = i, [str(c).strip() if c is not None else "" for c in r]
            break
    if head_i is None:
        raise ValueError(
            "Couldn't find the heading row — expected a row with 'Payroll ID' "
            "and 'Employee' on it. Is this the Standard Payroll Summary Report?")

    col = {}
    for j, h in enumerate(head):
        if h:
            col.setdefault(_norm(h), j)

    def find(names):
        return [col[_norm(n)] for n in names if _norm(n) in col]

    missing = [k for k, names in COST_PARTS.items() if not find(names)]
    if missing:
        warnings.append(
            "These columns weren't on the report, so they count as zero: "
            + ", ".join(COST_PARTS[k][0] for k in missing))

    # --- the period, from the lines above the heading ---
    start = end = None
    for r in rows[:head_i]:
        for c in (r or []):
            if c and re.search(r"20\d{2}", str(c)) and re.search(r"[A-Za-z]{3}", str(c)):
                s, e = parse_period(c)
                if s and e:
                    start, end = s, e
                    break
        if start:
            break
    if not start:
        warnings.append("Couldn't read the pay period off the report — "
                        "it will need setting by hand.")

    i_id, i_nm = col.get("payrollid"), col.get("employee")
    i_dept = col.get("department")

    def cell(r, j):
        return r[j] if j is not None and j < len(r) else None

    people = []
    for r in rows[head_i + 1:]:
        if not r:
            continue
        pid = str(cell(r, i_id) or "").strip()
        nm = str(cell(r, i_nm) or "").strip()
        if not pid or not nm:
            continue                     # the totals row has neither

        vals = {}
        for key, names in {**COST_PARTS, **INFO_PARTS}.items():
            vals[key] = round(sum(_num(cell(r, j)) for j in find(names)), 2)

        cost = round(vals["gross"] + vals["employer_nic"] + vals["employer_pension"], 2)
        people.append({
            "payroll_id": pid,
            "payroll_name": nm,
            "name": flip_name(nm),
            "payroll_dept": str(cell(r, i_dept) or "").strip(),
            **vals,
            "cost": cost,
        })

    if any(p["employer_pension_gross"] for p in people):
        warnings.append(
            "Some rows have an Employer GROSS Pension figure. Only the NET "
            "employer pension is counted, per the agreed definition — worth "
            "checking that's right for a salary-sacrifice scheme.")

    return people, start, end, warnings


def flip_name(s):
    """'Iovu, Florin' -> 'Florin Iovu'. Payroll writes surname first."""
    s = (s or "").strip()
    if "," in s:
        a, b = s.split(",", 1)
        return f"{b.strip()} {a.strip()}".strip()
    return s


def base_id(n):
    """'1122a' -> '1122'. Payroll adds a letter when someone is re-hired;
    Go2Clock and the app usually hold the bare number."""
    return re.sub(r"[^0-9]", "", str(n or ""))


def clean_name(s):
    """Same squashing as the Go2Clock matcher, so the two agree."""
    s = re.sub(r"\*+", "", s or "")
    s = re.sub(r"\([^)]*\)", "", s)
    return re.sub(r"[^a-z]", "", s.lower())


def match_payroll_person(person, timesheet_index, staff_index):
    """Tie a payroll row to a Go2Clock person and a member of staff.

    Payroll ID first — it agrees with Go2Clock 70 times out of 82, which is
    better than any name comparison. Then the bare number, then the name.
    A name that matches more than one person is reported, never guessed.
    """
    pid = (person.get("payroll_id") or "").strip()
    bare = base_id(pid)
    nm = person.get("name") or ""

    for key, how in ((pid, "payroll number"), (bare, "payroll number")):
        if key and key in timesheet_index["key"]:
            hit = timesheet_index["key"][key]
            return {"key": hit["key"], "employee_number": hit.get("employee_number"),
                    "staff_name": hit.get("staff_name"),
                    "matched_by": how, "confidence": "high"}

    hits = timesheet_index["name"].get(clean_name(nm)) or []
    if len(hits) == 1:
        return {"key": hits[0]["key"], "employee_number": hits[0].get("employee_number"),
                "staff_name": hits[0].get("staff_name"),
                "matched_by": "name", "confidence": "medium"}
    if len(hits) > 1:
        return {"key": None, "employee_number": None, "staff_name": None,
                "matched_by": "name", "confidence": "ambiguous",
                "note": f"{len(hits)} people on the clock share that name"}

    # not on the clock at all — still try the Name List, so salaried office
    # staff who don't clock in are costed
    shits = staff_index["name"].get(clean_name(nm)) or []
    if bare and bare in staff_index["num"]:
        s = staff_index["num"][bare]
        return {"key": None, "employee_number": s["employee_number"],
                "staff_name": s["name"], "matched_by": "Name List number",
                "confidence": "medium",
                "note": "No clock record — hours can't be worked out"}
    if len(shits) == 1:
        return {"key": None, "employee_number": shits[0]["employee_number"],
                "staff_name": shits[0]["name"], "matched_by": "Name List name",
                "confidence": "medium",
                "note": "No clock record — hours can't be worked out"}

    return {"key": None, "employee_number": None, "staff_name": None,
            "matched_by": None, "confidence": "none",
            "note": "Not on the clock or the Name List — a contractor, "
                    "a director, or someone who needs adding"}
