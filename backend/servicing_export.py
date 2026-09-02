"""Service-manager friendly Excel workbook for Pre Service Checks + Workshop Service records."""
import io
from typing import List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REPAIR_FILL = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")
ORDER_FILL = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
OTHER_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WRAP = Alignment(wrap_text=True, vertical="top")

RECORD_LABELS = {"pre_service_check": "Pre Service Check", "workshop_service": "Workshop Service"}
STATUS_LABELS = {"satisfactory": "OK", "unsatisfactory": "NEEDS WORK", "n/a": "N/A", "unchecked": ""}


def _date_parts(completed_at) -> Tuple[str, str]:
    s = str(completed_at or "")
    return (s[:10], s[11:16] if len(s) > 16 else "")


def _machine(c: dict) -> str:
    return f"{c.get('machine_make', '')} {c.get('machine_model', '')}".strip()


def _write_sheet(ws, headers: List[str], widths: List[int], rows: List[list], fills=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for idx, row in enumerate(rows):
        ws.append(row)
        fill = fills[idx] if fills else None
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
            if fill:
                cell.fill = fill
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def action_rows(checklists: List[dict]) -> Tuple[List[list], List]:
    """One row per thing the service manager must repair / order / look at."""
    rows, fills = [], []
    for c in checklists:
        date_str, _ = _date_parts(c.get("completed_at"))
        base = [date_str, _machine(c)]
        tail = [c.get("staff_name", ""), RECORD_LABELS.get(c.get("check_type"), c.get("check_type", "")), ""]
        if c.get("check_type") == "pre_service_check":
            for item in c.get("checklist_items") or []:
                if item.get("status") == "unsatisfactory":
                    rows.append(base + ["Repair", item.get("item", ""), item.get("notes") or ""] + tail)
                    fills.append(REPAIR_FILL)
            for part in c.get("parts_required") or []:
                rows.append(base + ["Order Part", "Parts required", part] + tail)
                fills.append(ORDER_FILL)
            if c.get("workshop_notes"):
                rows.append(base + ["Other Issue", "Any other Parts or Issues", c["workshop_notes"]] + tail)
                fills.append(OTHER_FILL)
        elif c.get("check_type") == "workshop_service":
            rows.append(base + ["Workshop Service", "Work completed", c.get("workshop_notes") or ""] + tail)
            fills.append(None)
    return rows, fills


def parts_rows(checklists: List[dict]) -> List[list]:
    rows = []
    for c in checklists:
        date_str, _ = _date_parts(c.get("completed_at"))
        for part in c.get("parts_required") or []:
            rows.append([part, _machine(c), date_str, c.get("staff_name", ""), ""])
    return rows


def section_columns(checklists: List[dict]) -> List[str]:
    seen = []
    for c in checklists:
        if c.get("check_type") != "pre_service_check":
            continue
        for item in c.get("checklist_items") or []:
            name = item.get("item", "")
            if name and name not in seen:
                seen.append(name)
    return seen


def section_text(item: dict) -> str:
    label = STATUS_LABELS.get(item.get("status", ""), item.get("status", ""))
    notes = (item.get("notes") or "").strip()
    return f"{label} - {notes}" if notes else label


def sheet_rows(checklists: List[dict], sections: List[str]) -> List[list]:
    rows = []
    for c in checklists:
        date_str, time_str = _date_parts(c.get("completed_at"))
        items = {i.get("item", ""): i for i in (c.get("checklist_items") or [])}
        needs_work = sum(1 for i in items.values() if i.get("status") == "unsatisfactory")
        row = [date_str, time_str, c.get("machine_make", ""), c.get("machine_model", ""), c.get("staff_name", ""),
               RECORD_LABELS.get(c.get("check_type"), c.get("check_type", ""))]
        row += [section_text(items[s]) if s in items else "" for s in sections]
        row += [c.get("workshop_notes") or "", "; ".join(c.get("parts_required") or []), needs_work, c.get("id", "")]
        rows.append(row)
    return rows


def build_servicing_workbook(checklists: List[dict]) -> io.BytesIO:
    wb = Workbook()

    ws = wb.active
    ws.title = "Action List"
    rows, fills = action_rows(checklists)
    _write_sheet(ws, ["Date", "Machine", "Action", "Section / Part", "Details", "Reported By", "Record Type", "Done?"],
                 [12, 28, 16, 30, 60, 18, 18, 8], rows, fills)

    ws_parts = wb.create_sheet("Parts to Order")
    _write_sheet(ws_parts, ["Part", "Machine", "Date", "Reported By", "Ordered?"], [40, 28, 12, 18, 10], parts_rows(checklists))

    ws_sheets = wb.create_sheet("Service Sheets")
    sections = section_columns(checklists)
    headers = ["Date", "Time", "Machine Make", "Machine Model", "Staff", "Record Type"] + sections + \
              ["Any other Parts or Issues", "Parts Required", "Needs Work Count", "Record ID"]
    widths = [12, 8, 16, 22, 18, 18] + [28] * len(sections) + [45, 40, 10, 38]
    srows = sheet_rows(checklists, sections)
    _write_sheet(ws_sheets, headers, widths, srows)
    first_section_col = 7
    for r_idx, row in enumerate(srows, 2):
        for s_idx in range(len(sections)):
            value = row[first_section_col - 1 + s_idx]
            if isinstance(value, str) and value.startswith("NEEDS WORK"):
                ws_sheets.cell(row=r_idx, column=first_section_col + s_idx).fill = REPAIR_FILL

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
