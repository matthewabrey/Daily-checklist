"""Backend tests for the Checks / Servicing category filter on checklists + exports"""
import csv
import io
import os

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

SERVICING = {"pre_service_check", "workshop_service"}
EXPORT_HEADERS = ["ID", "Staff Name", "Machine Make", "Machine Model", "Check Type", "Completed At",
                  "Status", "Satisfactory", "Unsatisfactory", "Total", "Needs Work / Repairs", "Notes",
                  "Workshop Details", "Parts Required"]
SECTIONS = [
    "Gun Carriage General", "Gun Carriage Wheels and Axles", "Gun", "Hydraulic Rams",
    "Drum", "Guards", "Computer/Computer Box",
]


@pytest.fixture(scope="module")
def api():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def created_ids():
    return []


@pytest.fixture(scope="module", autouse=True)
def cleanup(created_ids):
    yield
    if created_ids:
        print("\nNo DELETE endpoint; TEST_ records left in DB: " + ", ".join(created_ids))


# --- Module: GET /api/checklists?category= ---
class TestChecklistsCategory:
    def test_servicing_only_returns_servicing_types(self, api):
        r = api.get(f"{BASE_URL}/api/checklists", params={"limit": 100, "category": "servicing"}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        data = r.json()
        assert isinstance(data, list)
        types = {c["check_type"] for c in data}
        assert types <= SERVICING, f"Unexpected types in servicing: {types - SERVICING}"
        assert len(data) > 0, "Expected some servicing records in DB"

    def test_checks_excludes_servicing_types(self, api):
        r = api.get(f"{BASE_URL}/api/checklists", params={"limit": 200, "category": "checks"}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        data = r.json()
        types = {c["check_type"] for c in data}
        assert not (types & SERVICING), f"Servicing types leaked into checks: {types & SERVICING}"
        assert len(data) > 0

    def test_no_category_is_backwards_compatible(self, api):
        r = api.get(f"{BASE_URL}/api/checklists", params={"limit": 20}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        data = r.json()
        assert len(data) > 0
        assert all("check_type" in c and "id" in c for c in data)
        assert all("_id" not in c for c in data)

    def test_category_plus_check_type(self, api):
        r = api.get(f"{BASE_URL}/api/checklists",
                    params={"limit": 50, "category": "servicing", "check_type": "workshop_service"}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        data = r.json()
        assert len(data) > 0
        assert all(c["check_type"] == "workshop_service" for c in data)

    def test_invalid_category_returns_all(self, api):
        r = api.get(f"{BASE_URL}/api/checklists", params={"limit": 20, "category": "bogus"}, timeout=60)
        assert r.status_code == 200
        assert len(r.json()) > 0


# --- Module: GET /api/checklists/today?category= ---
class TestTodayCategory:
    def test_today_servicing(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/today", params={"category": "servicing"}, timeout=60)
        assert r.status_code == 200, r.text[:400]
        data = r.json()
        assert isinstance(data, list)
        assert all(c["check_type"] in SERVICING for c in data)

    def test_today_no_param(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/today", timeout=60)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_today_checks(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/today", params={"category": "checks"}, timeout=60)
        assert r.status_code == 200
        assert all(c["check_type"] not in SERVICING for c in r.json())


def _parse_csv(resp):
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
    return rows[0], rows[1:]


# --- Module: create a pre_service_check record used by export assertions ---
class TestCreatePreServiceRecord:
    def test_create(self, api, created_ids):
        statuses = ["satisfactory"] * 4 + ["unsatisfactory", "satisfactory", "n/a"]
        items = [
            {"item": s, "status": statuses[i],
             "notes": "TEST_worn seal" if statuses[i] == "unsatisfactory" else ""}
            for i, s in enumerate(SECTIONS)
        ]
        body = {
            "employee_number": "4444",
            "staff_name": "TEST_Agent",
            "machine_make": "Perrot",
            "machine_model": "Machine 2",
            "check_type": "pre_service_check",
            "checklist_items": items,
            "workshop_notes": "TEST_other issue",
            "parts_required": ["TEST_part A", "TEST_part B"],
        }
        r = api.post(f"{BASE_URL}/api/checklists", json=body, timeout=60)
        assert r.status_code == 200, r.text[:500]
        d = r.json()
        assert d["check_type"] == "pre_service_check"
        created_ids.append(d["id"])
        # verify persisted
        g = api.get(f"{BASE_URL}/api/checklists/{d['id']}", timeout=30)
        assert g.status_code == 200
        assert g.json()["parts_required"] == ["TEST_part A", "TEST_part B"]


# --- Module: CSV export with category ---
class TestCsvExport:
    def test_csv_servicing(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", params={"category": "servicing"}, timeout=120)
        assert r.status_code == 200, r.text[:300]
        assert "all_servicing.csv" in r.headers.get("content-disposition", "")
        header, rows = _parse_csv(r)
        assert header == EXPORT_HEADERS
        assert len(header) == 14
        assert len(rows) > 0
        bad = {row[4] for row in rows} - SERVICING
        assert not bad, f"Non-servicing rows in servicing export: {bad}"

    def test_csv_no_param(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", timeout=120)
        assert r.status_code == 200
        assert "all_checks.csv" in r.headers.get("content-disposition", "")
        header, rows = _parse_csv(r)
        assert header == EXPORT_HEADERS
        assert "daily_check" in {row[4] for row in rows}

    def test_csv_checks_has_no_servicing(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", params={"category": "checks"}, timeout=120)
        assert r.status_code == 200
        assert "all_checks.csv" in r.headers.get("content-disposition", "")
        _, rows = _parse_csv(r)
        assert not ({row[4] for row in rows} & SERVICING)

    def test_csv_pre_service_row_values(self, api, created_ids):
        cid = created_ids[0]
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", params={"category": "servicing"}, timeout=120)
        _, rows = _parse_csv(r)
        match = [row for row in rows if row[0] == cid]
        assert match, f"Created record {cid} not present in servicing CSV export"
        row = match[0]
        assert row[1] == "TEST_Agent"
        assert row[4] == "pre_service_check"
        assert row[7] == "5", f"Satisfactory expected 5, got {row[7]}"
        assert row[8] == "1", f"Unsatisfactory expected 1, got {row[8]}"
        assert row[9] == "7", f"Total expected 7, got {row[9]}"
        assert "Drum: TEST_worn seal" in row[10], f"Needs Work column: {row[10]!r}"
        assert "Drum: TEST_worn seal" in row[11], f"Notes column: {row[11]!r}"
        assert row[12] == "TEST_other issue"
        assert row[13] == "TEST_part A; TEST_part B"


# --- Module: Excel export with category ---
class TestExcelExport:
    def test_excel_servicing(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel", params={"category": "servicing"}, timeout=180)
        assert r.status_code == 200, r.text[:300]
        assert "all_servicing.xlsx" in r.headers.get("content-disposition", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Action List", "Parts to Order", "Service Sheets"]
        ws = wb["Service Sheets"]
        header = [c.value for c in ws[1]]
        assert header[:6] == ["Date", "Time", "Machine Make", "Machine Model", "Staff", "Record Type"]
        types = {ws.cell(row=i, column=6).value for i in range(2, ws.max_row + 1)}
        assert types <= {"Pre Service Check", "Workshop Service"}, f"Unexpected: {types}"

    def test_excel_no_param(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel", timeout=180)
        assert r.status_code == 200
        assert "all_checks.xlsx" in r.headers.get("content-disposition", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.active.title == "All Checks"
        assert [c.value for c in wb.active[1]] == EXPORT_HEADERS
