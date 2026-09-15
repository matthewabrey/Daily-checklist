"""Iteration 16 - new feature: AssetList-driven service check templates + generic pre-service."""
import os
import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
BASE_URL = base_url.rstrip("/")

EXPECTED_SECTIONS = [
    "Gun Carriage General", "Gun Carriage Wheels and Axles", "Gun and Nozel",
    "Hydraulic Rams", "Drum and Bearings", "Guards", "Computer/Computer Box",
    "Pipe", "Solar Panel", "Raindancer", "Raindancer Serial Number", "Seals",
    "Intake Pipe Work",
]

ASSET_FILE = "/tmp/AssetList2.xlsx"


@pytest.fixture(scope="module")
def api():
    s = requests.Session()
    return s


# --- Module: service-check-templates/for-asset ---
class TestForAssetEndpoint:
    def test_irrigator_perrot(self, api):
        r = api.get(f"{BASE_URL}/api/service-check-templates/for-asset",
                    params={"check_type": "Irrigator", "make": "Perrot"}, timeout=30)
        assert r.status_code == 200
        d = r.json()
        assert d["check_type"] == "Irrigator"
        assert d["name"] == "Irrigator Pre Service Check"
        assert d["sections"] == EXPECTED_SECTIONS
        assert len(d["sections"]) == 13

    @pytest.mark.parametrize("make", ["Perrot", "Briggs", "Bauer"])
    def test_case_insensitive_all_makes(self, api, make):
        r = api.get(f"{BASE_URL}/api/service-check-templates/for-asset",
                    params={"check_type": "irrigator", "make": make}, timeout=30)
        assert r.status_code == 200, f"{make}: {r.text[:200]}"
        assert r.json()["check_type"] == "Irrigator"

    def test_tractor_jcb_returns_404(self, api):
        r = api.get(f"{BASE_URL}/api/service-check-templates/for-asset",
                    params={"check_type": "Tractor", "make": "JCB"}, timeout=30)
        assert r.status_code == 404

    def test_no_params_returns_404(self, api):
        r = api.get(f"{BASE_URL}/api/service-check-templates/for-asset", timeout=30)
        assert r.status_code == 404

    def test_list_exactly_one_irrigator_no_legacy(self, api):
        r = api.get(f"{BASE_URL}/api/service-check-templates", timeout=30)
        assert r.status_code == 200
        data = r.json()
        irrigators = [t for t in data if t.get("check_type") == "Irrigator"]
        assert len(irrigators) == 1
        # legacy make-only (check_type null) should NOT exist
        legacy = [t for t in data if not t.get("check_type")]
        assert legacy == [], f"Legacy templates present: {legacy}"


# --- Module: admin upload assets + idempotency + tab mapping ---
class TestAssetsUpload:
    def _upload(self, api):
        with open(ASSET_FILE, "rb") as fh:
            files = {"file": ("AssetList2.xlsx", fh,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
            return api.post(f"{BASE_URL}/api/admin/upload-assets-file",
                            files=files, timeout=180)

    def test_first_upload_message_and_sheets(self, api):
        r = self._upload(api)
        assert r.status_code == 200, r.text[:500]
        d = r.json()
        assert "17 daily check lists and 1 pre service sheets" in d.get("message", ""), d.get("message")
        assert d.get("service_templates_created") == 1
        sheets = d.get("processed_sheets", [])
        sheets_str = " || ".join(sheets)
        assert "Irrigators - Pre Service Sheet" in sheets_str and "Pre Service Sheet: 13 sections" in sheets_str
        assert "Irrigators - Daily Check" in sheets_str and "5 items" in sheets_str and "5 compulsory" in sheets_str
        # Tractor Hire must map to Tractor Hire, not Tractor
        assert any("Tractor Hire -> Tractor Hire" in s for s in sheets), sheets_str

    def test_second_upload_idempotent(self, api):
        r = self._upload(api)
        assert r.status_code == 200
        d = r.json()
        assert d.get("added") == 0, d
        r2 = api.get(f"{BASE_URL}/api/service-check-templates", timeout=30)
        irrigators = [t for t in r2.json() if t["check_type"] == "Irrigator"]
        assert len(irrigators) == 1


# --- Module: daily checklist templates regression ---
class TestDailyTemplates:
    def test_irrigator_daily_5_items_all_compulsory(self, api):
        r = api.get(f"{BASE_URL}/api/checklist-templates/Irrigator", timeout=30)
        assert r.status_code == 200
        body = r.json()
        items = body.get("items") if isinstance(body, dict) else body
        assert len(items) == 5, items
        assert all(i.get("compulsory") is True for i in items), items
        texts = [i["item"] for i in items]
        assert "Check all Lights are Working" in texts[0]
        assert any("safe for road use" in t.lower() for t in texts)

    def test_tractor_daily_18_items(self, api):
        r = api.get(f"{BASE_URL}/api/checklist-templates/Tractor", timeout=30)
        assert r.status_code == 200
        body = r.json()
        items = body.get("items") if isinstance(body, dict) else body
        assert len(items) == 18, len(items)


# --- Module: diagnostics ---
class TestTemplateDiagnostics:
    def test_diagnostics_service_templates(self, api):
        r = api.get(f"{BASE_URL}/api/admin/template-diagnostics", timeout=60)
        assert r.status_code == 200
        d = r.json()
        st = d.get("service_templates", [])
        irr = next((t for t in st if t["check_type"] == "Irrigator"), None)
        assert irr is not None, st
        assert irr["section_count"] == 13
        assert irr["assets_using_this"] >= 90
        assert irr.get("sheet_name")
        assert d.get("missing_templates", []) == []


# --- Module: pre_service_check submissions (Irrigator + Generic) ---
class TestPreServiceSubmissions:
    def test_irrigator_pre_service_submit_and_get(self, api):
        items = [{"item": s, "status": "satisfactory", "notes": ""} for s in EXPECTED_SECTIONS]
        items[3] = {"item": EXPECTED_SECTIONS[3], "status": "unsatisfactory", "notes": "TEST_ram leak"}
        body = {
            "employee_number": "4444", "staff_name": "Admin User",
            "machine_make": "Perrot", "machine_model": "Machine 1",
            "check_type": "pre_service_check",
            "checklist_items": items,
            "workshop_notes": "TEST_iter16 irrigator",
            "parts_required": ["TEST_Seal kit 998-123", "TEST_Hose 12mm"],
        }
        r = api.post(f"{BASE_URL}/api/checklists", json=body, timeout=60)
        assert r.status_code == 200, r.text[:400]
        cid = r.json()["id"]
        g = api.get(f"{BASE_URL}/api/checklists/{cid}", timeout=30).json()
        assert g["parts_required"] == body["parts_required"]
        assert len(g["checklist_items"]) == 13
        assert g["checklist_items"][3]["status"] == "unsatisfactory"

    def test_generic_pre_service_2_items(self, api):
        body = {
            "employee_number": "4444", "staff_name": "Admin User",
            "machine_make": "JCB", "machine_model": "TEST_Fastrac 4220",
            "check_type": "pre_service_check",
            "checklist_items": [
                {"item": "Gearbox", "status": "unsatisfactory", "notes": "TEST_whine on load"},
                {"item": "Hydraulic hoses", "status": "satisfactory", "notes": ""},
            ],
            "workshop_notes": "TEST_iter16 generic",
            "parts_required": ["TEST_Seal kit 998-123"],
        }
        r = api.post(f"{BASE_URL}/api/checklists", json=body, timeout=60)
        assert r.status_code == 200, r.text[:400]
        d = r.json()
        assert d["check_type"] == "pre_service_check"
        assert len(d["checklist_items"]) == 2
        cid = d["id"]
        g = api.get(f"{BASE_URL}/api/checklists/{cid}", timeout=30).json()
        assert g["parts_required"] == ["TEST_Seal kit 998-123"]
        assert g["checklist_items"][0]["item"] == "Gearbox"


# --- Module: servicing Excel export regression with 13 section columns ---
class TestServicingExcelExport:
    def test_servicing_excel_3_sheets_with_13_sections(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel",
                    params={"category": "servicing"}, timeout=120)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml"
        ), r.headers.get("content-type")
        from io import BytesIO
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(r.content))
        assert len(wb.sheetnames) == 3, wb.sheetnames
        # find service sheets tab
        svc_sheet = next((n for n in wb.sheetnames if "service" in n.lower()), None)
        assert svc_sheet, wb.sheetnames
        ws = wb[svc_sheet]
        header = [c.value for c in ws[1]]
        for sec in EXPECTED_SECTIONS:
            assert sec in header, f"missing section col {sec!r} in {header}"
