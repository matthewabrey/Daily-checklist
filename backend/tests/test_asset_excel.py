"""Unit tests for asset_excel.py - AssetList.xlsx tab naming convention"""
import sys
from io import BytesIO

import openpyxl

sys.path.insert(0, "/app/backend")
from asset_excel import parse_asset_workbook, resolve_sheet_check_types, split_sheet_name  # noqa: E402


def build_workbook(front_rows, tabs):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Front Sheet"
    ws.append(["Check Type", "NameColumn", "MakeColumn", "Crop", "Year"])
    for r in front_rows:
        ws.append(r)
    for name, rows in tabs.items():
        t = wb.create_sheet(title=name)
        for r in rows:
            t.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_split_sheet_name_variants():
    assert split_sheet_name("Irrigators - Pre Service Sheet ") == ("Irrigators", "service")
    assert split_sheet_name("Irrigators - Service Sheet") == ("Irrigators", "service")
    assert split_sheet_name("Forklift - Pre Service Check") == ("Forklift", "service")
    assert split_sheet_name("Irrigators - Daily Check") == ("Irrigators", "daily")
    assert split_sheet_name("Forklift") == ("Forklift", "daily")
    assert split_sheet_name("MEWP - GENIE etc") == ("MEWP - GENIE etc", "daily")
    assert split_sheet_name("Grader Checklist") == ("Grader Checklist", "daily")


def test_resolve_partial_never_overwrites_exact():
    out = resolve_sheet_check_types(
        ["Irrigators - Pre Service Sheet ", "Irrigators - Daily Check", "Tractor Hire", "Tractor", "AD Plant Monthly Inspections"],
        ["Irrigator", "Tractor", "Monthly Inspections"],
    )
    assert out == [
        ("Irrigators - Pre Service Sheet ", "Irrigator", "service"),
        ("Irrigators - Daily Check", "Irrigator", "daily"),
        ("Tractor Hire", "Tractor Hire", "daily"),
        ("Tractor", "Tractor", "daily"),
        ("AD Plant Monthly Inspections", "Monthly Inspections", "daily"),
    ]


def test_parse_workbook_daily_and_service_tabs():
    content = build_workbook(
        [["Irrigator", "1 AF 450", "Perrot"], ["Irrigator", "B1 R64", "Briggs"], ["Tractor", "Fendt 720", "Fendt"]],
        {
            "Irrigators - Pre Service Sheet ": [[None, None], [1, "Gun Carriage General"], [2, "Gun and Nozel"], [3, "Pipe"]],
            "Irrigators - Daily Check": [["Item", "Compulsary"], ["Check all Lights are Working", "yes"], ["Check Tyre Pressure", "no"]],
            "Tractor": [["Item", "Compulsary"], ["Walk-Around Inspection", "yes"]],
            "Empty Tab": [],
        },
    )
    parsed = parse_asset_workbook(content)
    assert len(parsed["assets"]) == 3
    assert [t["check_type"] for t in parsed["service_templates"]] == ["Irrigator"]
    svc = parsed["service_templates"][0]
    assert svc["sections"] == ["Gun Carriage General", "Gun and Nozel", "Pipe"]
    assert svc["name"] == "Irrigator Pre Service Check"
    assert svc["sheet_name"] == "Irrigators - Pre Service Sheet "
    daily = {t["check_type"]: t for t in parsed["checklist_templates"]}
    assert set(daily) == {"Irrigator", "Tractor"}
    assert daily["Irrigator"]["items"][0] == {"item": "Check all Lights are Working", "compulsory": True, "critical": False, "photo_required": False}
    assert daily["Irrigator"]["items"][1]["compulsory"] is False
    assert any("Empty Tab -> skipped" in s for s in parsed["processed_sheets"])
    assert any("Pre Service Sheet: 3 sections" in s for s in parsed["processed_sheets"])


def test_missing_columns_raises():
    wb = openpyxl.Workbook()
    wb.active.append(["Foo", "Bar"])
    buf = BytesIO()
    wb.save(buf)
    try:
        parse_asset_workbook(buf.getvalue())
        assert False, "expected ValueError"
    except ValueError as e:
        assert "Check Type" in str(e)
