"""Backend tests for the service-manager servicing export workbook + notes-in-export bug fix."""
import csv
import io
import os

import pytest
import requests
from dotenv import dotenv_values
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

EXPORT_HEADERS = ["ID", "Staff Name", "Machine Make", "Machine Model", "Check Type", "Completed At",
                  "Status", "Satisfactory", "Unsatisfactory", "Total", "Needs Work / Repairs", "Notes",
                  "Workshop Details", "Parts Required"]


@pytest.fixture(scope="module")
def api():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def state():
    return {}


@pytest.fixture(scope="module", autouse=True)
def cleanup(state):
    yield
    ids = [v for k, v in state.items() if k.endswith("_id")]
    if ids:
        print("\nNo DELETE endpoint; TEST_ records left in DB: " + ", ".join(ids))


# --- Module: seed records (pre_service_check with notes on OK item + daily_check) ---
class TestSeed:
    def test_create_pre_service_check(self, api, state):
        items = [
            {"item": "Gun Carriage General", "status": "satisfactory", "notes": "TEST_slight rust ok"},
            {"item": "Gun Carriage Wheels and Axles", "status": "satisfactory", "notes": ""},
            {"item": "Gun", "status": "satisfactory", "notes": ""},
            {"item": "Hydraulic Rams", "status": "unsatisfactory", "notes": "TEST_left ram leaking"},
            {"item": "Drum", "status": "unsatisfactory", "notes": "TEST_drum bearing worn"},
            {"item": "Guards", "status": "satisfactory", "notes": ""},
            {"item": "Computer/Computer Box", "status": "n/a", "notes": ""},
        ]
        body = {
            "employee_number": "4444",
            "staff_name": "TEST_Notes",
            "machine_make": "Perrot",
            "machine_model": "Machine 4",
            "check_type": "pre_service_check",
            "checklist_items": items,
            "workshop_notes": "TEST_hose split near gun",
            "parts_required": ["TEST_Drum bearing", "TEST_Ram seal kit"],
        }
        r = api.post(f"{BASE_URL}/api/checklists", json=body, timeout=60)
        assert r.status_code == 200, r.text[:500]
        d = r.json()
        state["pre_service_id"] = d["id"]
        g = api.get(f"{BASE_URL}/api/checklists/{d['id']}", timeout=30)
        assert g.status_code == 200
        gd = g.json()
        assert gd["machine_make"] == "Perrot" and gd["machine_model"] == "Machine 4"
        assert gd["parts_required"] == ["TEST_Drum bearing", "TEST_Ram seal kit"]
        assert gd["workshop_notes"] == "TEST_hose split near gun"
        by_item = {i["item"]: i for i in gd["checklist_items"]}
        assert by_item["Gun Carriage General"]["notes"] == "TEST_slight rust ok"

    def test_create_daily_check(self, api, state):
        body = {
            "employee_number": "4444",
            "staff_name": "TEST_Notes",
            "machine_make": "JCB",
            "machine_model": "TEST_Daily Machine",
            "check_type": "daily_check",
            "checklist_items": [{"item": "Oil level", "status": "satisfactory", "notes": "TEST_daily note"}],
        }
        r = api.post(f"{BASE_URL}/api/checklists", json=body, timeout=60)
        assert r.status_code == 200, r.text[:500]
        state["daily_id"] = r.json()["id"]


def _servicing_wb(api):
    r = api.get(f"{BASE_URL}/api/checklists/export/excel", params={"category": "servicing"}, timeout=180)
    assert r.status_code == 200, r.text[:300]
    return r, load_workbook(io.BytesIO(r.content))


def _rows_as_dicts(ws):
    header = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        out.append({header[i]: c for i, c in enumerate(row) if i < len(header)})
    return header, out


# --- Module: Service Sheets sheet (notes for every section exported) ---
class TestServiceSheets:
    def test_sheetnames(self, api):
        r, wb = _servicing_wb(api)
        assert "all_servicing.xlsx" in r.headers.get("content-disposition", "")
        assert wb.sheetnames == ["Action List", "Parts to Order", "Service Sheets"]

    def test_service_sheet_row_values(self, api, state):
        _, wb = _servicing_wb(api)
        ws = wb["Service Sheets"]
        header, rows = _rows_as_dicts(ws)
        assert header[:6] == ["Date", "Time", "Machine Make", "Machine Model", "Staff", "Record Type"]
        for tail in ["Any other Parts or Issues", "Parts Required", "Needs Work Count", "Record ID"]:
            assert tail in header
        match = [r for r in rows if r["Record ID"].value == state["pre_service_id"]]
        assert match, f"record {state['pre_service_id']} missing from Service Sheets"
        row = match[0]
        assert row["Machine Make"].value == "Perrot"
        assert row["Machine Model"].value == "Machine 4"
        assert row["Staff"].value == "TEST_Notes"
        assert row["Record Type"].value == "Pre Service Check"
        assert row["Gun Carriage General"].value == "OK - TEST_slight rust ok"
        assert row["Hydraulic Rams"].value == "NEEDS WORK - TEST_left ram leaking"
        assert row["Drum"].value == "NEEDS WORK - TEST_drum bearing worn"
        assert row["Computer/Computer Box"].value == "N/A"
        assert row["Gun"].value == "OK"
        assert row["Any other Parts or Issues"].value == "TEST_hose split near gun"
        assert row["Parts Required"].value == "TEST_Drum bearing; TEST_Ram seal kit"
        assert row["Needs Work Count"].value == 2

    def test_needs_work_cells_red_fill(self, api, state):
        _, wb = _servicing_wb(api)
        _, rows = _rows_as_dicts(wb["Service Sheets"])
        row = [r for r in rows if r["Record ID"].value == state["pre_service_id"]][0]
        for col in ["Hydraulic Rams", "Drum"]:
            rgb = row[col].fill.start_color.rgb
            assert rgb and "FDE2E2" in str(rgb), f"{col} fill {rgb}"


# --- Module: Action List sheet ---
class TestActionList:
    def test_header_and_rows(self, api):
        _, wb = _servicing_wb(api)
        ws = wb["Action List"]
        header = [c.value for c in ws[1]]
        assert header == ["Date", "Machine", "Action", "Section / Part", "Details",
                          "Reported By", "Record Type", "Done?"]
        _, rows = _rows_as_dicts(ws)
        mine = [r for r in rows if r["Machine"].value == "Perrot Machine 4"
                and r["Reported By"].value == "TEST_Notes"]
        assert mine, "no Action List rows for Perrot Machine 4 / TEST_Notes"
        triples = {(r["Action"].value, r["Section / Part"].value, r["Details"].value) for r in mine}
        assert ("Repair", "Hydraulic Rams", "TEST_left ram leaking") in triples, triples
        assert ("Repair", "Drum", "TEST_drum bearing worn") in triples, triples
        assert ("Order Part", "Parts required", "TEST_Drum bearing") in triples, triples
        assert ("Order Part", "Parts required", "TEST_Ram seal kit") in triples, triples
        assert ("Other Issue", "Any other Parts or Issues", "TEST_hose split near gun") in triples, triples
        assert all(r["Record Type"].value == "Pre Service Check" for r in mine)

    def test_row_fills(self, api):
        _, wb = _servicing_wb(api)
        _, rows = _rows_as_dicts(wb["Action List"])
        mine = [r for r in rows if r["Machine"].value == "Perrot Machine 4"
                and r["Reported By"].value == "TEST_Notes"]
        expected = {"Repair": "FDE2E2", "Order Part": "EDE9FE", "Other Issue": "FEF3C7"}
        for r in mine:
            action = r["Action"].value
            if action in expected:
                rgb = str(r["Action"].fill.start_color.rgb)
                assert expected[action] in rgb, f"{action} fill {rgb}"

    def test_workshop_service_rows_present(self, api):
        _, wb = _servicing_wb(api)
        _, rows = _rows_as_dicts(wb["Action List"])
        ws_rows = [r for r in rows if r["Action"].value == "Workshop Service"]
        # workshop_service records exist in the DB (asserted by category tests)
        assert ws_rows, "no Workshop Service rows in Action List"
        assert all(r["Record Type"].value == "Workshop Service" for r in ws_rows)


# --- Module: Parts to Order sheet ---
class TestPartsToOrder:
    def test_parts_sheet(self, api):
        _, wb = _servicing_wb(api)
        ws = wb["Parts to Order"]
        header = [c.value for c in ws[1]]
        assert header == ["Part", "Machine", "Date", "Reported By", "Ordered?"]
        pairs = {(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2)}
        assert ("TEST_Drum bearing", "Perrot Machine 4") in pairs
        assert ("TEST_Ram seal kit", "Perrot Machine 4") in pairs


def _parse_csv(resp):
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
    return rows[0], rows[1:]


# --- Module: CSV exports include notes for every check type ---
class TestCsvNotes:
    def test_servicing_csv_row(self, api, state):
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", params={"category": "servicing"}, timeout=120)
        assert r.status_code == 200
        header, rows = _parse_csv(r)
        assert header == EXPORT_HEADERS and len(header) == 14
        match = [x for x in rows if x[0] == state["pre_service_id"]]
        assert match, "created record missing from servicing CSV"
        row = match[0]
        assert row[10] == "Hydraulic Rams: TEST_left ram leaking; Drum: TEST_drum bearing worn", row[10]
        assert "Gun Carriage General: TEST_slight rust ok" in row[11], row[11]
        assert "Drum: TEST_drum bearing worn" in row[11], row[11]
        assert row[12] == "TEST_hose split near gun"
        assert row[13] == "TEST_Drum bearing; TEST_Ram seal kit"

    def test_checks_csv_has_daily_notes(self, api, state):
        r = api.get(f"{BASE_URL}/api/checklists/export/csv", params={"category": "checks"}, timeout=120)
        assert r.status_code == 200
        header, rows = _parse_csv(r)
        assert header == EXPORT_HEADERS
        match = [x for x in rows if x[0] == state["daily_id"]]
        assert match, "created daily_check missing from checks CSV"
        assert "TEST_daily note" in match[0][11], match[0][11]


# --- Module: plain Excel export regression ---
class TestPlainExcel:
    def test_excel_no_category(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel", timeout=180)
        assert r.status_code == 200
        assert "all_checks.xlsx" in r.headers.get("content-disposition", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["All Checks"]
        assert [c.value for c in wb["All Checks"][1]] == EXPORT_HEADERS

    def test_excel_checks_category(self, api):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel", params={"category": "checks"}, timeout=180)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["All Checks"]
        ws = wb["All Checks"]
        assert [c.value for c in ws[1]] == EXPORT_HEADERS
        types = {ws.cell(row=i, column=5).value for i in range(2, ws.max_row + 1)}
        assert not (types & {"pre_service_check", "workshop_service"}), types


# --- Module: by-machine export full notes ---
class TestByMachineExport:
    def test_by_machine_notes_full(self, api, state):
        r = api.get(f"{BASE_URL}/api/checklists/export/excel-by-machine",
                    params={"make": "Perrot", "name": "Machine 4"}, timeout=180)
        assert r.status_code == 200, r.text[:300]
        wb = load_workbook(io.BytesIO(r.content))
        assert "pre_service_check" in wb.sheetnames, wb.sheetnames
        ws = wb["pre_service_check"]
        header = [c.value for c in ws[1]]
        notes_col = header.index("Notes") + 1
        notes = [ws.cell(row=i, column=notes_col).value or "" for i in range(2, ws.max_row + 1)]
        joined = "\n".join(notes)
        assert "Hydraulic Rams: TEST_left ram leaking" in joined, joined[:800]
        assert "Gun Carriage General: TEST_slight rust ok" in joined, joined[:800]
