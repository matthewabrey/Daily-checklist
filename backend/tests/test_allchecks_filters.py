"""Tests for /all-checks server-side filters (make/model/check_type) + filtered CSV/Excel exports"""
import csv
import io
import os
from datetime import datetime, timezone

import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

EXPECTED_HEADER_LEN = 14


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


def get_json(client, path, **params):
    r = client.get(f"{BASE_URL}{path}", params=params, timeout=120)
    assert r.status_code == 200, f"{path} {params} -> {r.status_code}: {r.text[:300]}"
    return r.json()


# ---------------- GET /api/checklists filters ----------------
class TestChecklistListFilters:
    def test_make_filter(self, client):
        data = get_json(client, "/api/checklists", limit=200, make="JCB")
        assert isinstance(data, list) and len(data) > 0, "expected some JCB checklists"
        assert all(c["machine_make"] == "JCB" for c in data), \
            f"non-JCB makes: {set(c['machine_make'] for c in data)}"

    def test_make_and_check_type_filter(self, client):
        data = get_json(client, "/api/checklists", limit=200, make="JCB", check_type="daily_check")
        assert len(data) > 0
        assert all(c["machine_make"] == "JCB" and c["check_type"] == "daily_check" for c in data)

    def test_make_and_model_filter(self, client):
        base = get_json(client, "/api/checklists", limit=200, make="JCB")
        model = base[0]["machine_model"]
        data = get_json(client, "/api/checklists", limit=200, make="JCB", model=model)
        assert len(data) > 0
        assert all(c["machine_make"] == "JCB" and c["machine_model"] == model for c in data)

    def test_servicing_category_with_make(self, client):
        data = get_json(client, "/api/checklists", limit=200, category="servicing", make="Perrot")
        assert len(data) >= 1
        assert all(c["machine_make"] == "Perrot" for c in data)
        assert all(c["check_type"] in ("pre_service_check", "workshop_service") for c in data)

    def test_unknown_make_returns_empty(self, client):
        data = get_json(client, "/api/checklists", limit=200, make="NoSuchMake")
        assert data == []

    # regression
    def test_no_filters_returns_mixed(self, client):
        data = get_json(client, "/api/checklists", limit=20)
        assert len(data) > 0

    def test_comma_separated_check_type(self, client):
        data = get_json(client, "/api/checklists", limit=100, check_type="daily_check,fuel_mileage")
        assert len(data) > 0
        assert all(c["check_type"] in ("daily_check", "fuel_mileage") for c in data)


# ---------------- CSV export with filters ----------------
def parse_csv(resp):
    rows = list(csv.reader(io.StringIO(resp.text)))
    return rows[0], rows[1:]


class TestCsvExportFilters:
    def test_csv_make_and_check_type(self, client):
        r = client.get(f"{BASE_URL}/api/checklists/export/csv",
                       params={"category": "checks", "make": "JCB", "check_type": "daily_check"}, timeout=180)
        assert r.status_code == 200
        header, rows = parse_csv(r)
        assert len(header) == EXPECTED_HEADER_LEN, header
        make_i, type_i = header.index("Machine Make"), header.index("Check Type")
        assert len(rows) > 0
        assert all(row[make_i] == "JCB" and row[type_i] == "daily_check" for row in rows)

        expected = get_json(client, "/api/checklists", limit=10000, category="checks",
                            make="JCB", check_type="daily_check")
        assert len(rows) == len(expected), f"csv rows {len(rows)} != api {len(expected)}"

    def test_csv_model_narrows(self, client):
        base = get_json(client, "/api/checklists", limit=200, category="checks", make="JCB", check_type="daily_check")
        model = base[0]["machine_model"]
        r = client.get(f"{BASE_URL}/api/checklists/export/csv",
                       params={"category": "checks", "make": "JCB", "check_type": "daily_check", "model": model},
                       timeout=180)
        assert r.status_code == 200
        header, rows = parse_csv(r)
        model_i = header.index("Machine Model")
        assert len(rows) > 0
        assert all(row[model_i] == model for row in rows)
        expected = get_json(client, "/api/checklists", limit=10000, category="checks",
                            make="JCB", check_type="daily_check", model=model)
        assert len(rows) == len(expected)

    def test_csv_unknown_make_header_only(self, client):
        r = client.get(f"{BASE_URL}/api/checklists/export/csv",
                       params={"category": "checks", "make": "NoSuchMake"}, timeout=120)
        assert r.status_code == 200
        header, rows = parse_csv(r)
        assert len(header) == EXPECTED_HEADER_LEN
        assert rows == []

    def test_csv_today_filter(self, client):
        r = client.get(f"{BASE_URL}/api/checklists/export/csv",
                       params={"category": "checks", "today": "true"}, timeout=120)
        assert r.status_code == 200
        header, rows = parse_csv(r)
        today = datetime.now(timezone.utc).date().isoformat()
        done_i = header.index("Completed At")
        assert all(row[done_i].startswith(today) for row in rows), \
            f"non-today rows: {[row[done_i] for row in rows[:5]]}"


# ---------------- Excel export with filters ----------------
class TestExcelExportFilters:
    def test_excel_checks_make(self, client):
        from openpyxl import load_workbook
        r = client.get(f"{BASE_URL}/api/checklists/export/excel",
                       params={"category": "checks", "make": "JCB"}, timeout=240)
        assert r.status_code == 200, r.text[:300]
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert wb.sheetnames == ["All Checks"], wb.sheetnames
        ws = wb["All Checks"]
        rows = list(ws.values)
        header = list(rows[0])
        assert len(header) == EXPECTED_HEADER_LEN
        make_i = header.index("Machine Make")
        assert len(rows) > 1
        assert all(row[make_i] == "JCB" for row in rows[1:])

    def test_excel_servicing_perrot_three_sheets(self, client):
        from openpyxl import load_workbook
        r = client.get(f"{BASE_URL}/api/checklists/export/excel",
                       params={"category": "servicing", "make": "Perrot"}, timeout=240)
        assert r.status_code == 200, r.text[:300]
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        assert wb.sheetnames == ["Action List", "Parts to Order", "Service Sheets"], wb.sheetnames
        rows = list(wb["Service Sheets"].values)
        header = list(rows[0])
        make_i = header.index("Machine Make")
        assert len(rows) > 1, "Service Sheets has no data rows"
        assert all(row[make_i] == "Perrot" for row in rows[1:])
