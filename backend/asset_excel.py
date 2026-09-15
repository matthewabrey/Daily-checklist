"""Parse AssetList.xlsx.

Front sheet = assets (Check Type / Name / Make). Every other tab is a template for a check type:
  "<Check Type>"                      -> Daily Check list (legacy, still supported)
  "<Check Type> - Daily Check"        -> Daily Check list
  "<Check Type> - Pre Service Sheet"  -> Pre Service Check sections (also "Service Sheet" / "Pre Service Check")
Tab names are matched to the Check Type column case-insensitively, ignoring punctuation and plurals.
"""
import re
import uuid
from io import BytesIO
from typing import Dict, List, Optional, Tuple

import openpyxl

SUFFIX_RE = re.compile(
    r"^(?P<base>.+?)\s*[-–:]\s*(?P<kind>pre[\s-]*service(?:\s*(?:sheet|check|checks))?|service\s*(?:sheet|check|checks)|daily\s*checks?|checklist)\s*$",
    re.IGNORECASE,
)
ITEM_HEADERS = ("item", "items", "section", "sections", "task", "tasks", "description", "check", "checks")
JUNK_ITEMS = {"item", "check", "task", "description", "checklist", "safety", "section"}


def clean_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower().replace("checklist", ""))


def split_sheet_name(sheet_name: str) -> Tuple[str, str]:
    """Return (check type base name, 'daily' | 'service')."""
    name = sheet_name.strip()
    m = SUFFIX_RE.match(name)
    if not m:
        return name, "daily"
    kind = "service" if "service" in m.group("kind").lower() else "daily"
    return m.group("base").strip(), kind


def match_check_type(base: str, check_types: List[str], partial: bool = False) -> Optional[str]:
    """Exact match ignoring case/punctuation/plurals; with partial=True also allow containment."""
    target = clean_name(base)
    if not target:
        return None
    by_clean = {clean_name(ct): ct for ct in check_types}
    if target in by_clean:
        return by_clean[target]
    singular = target.rstrip("s")
    for cleaned, ct in by_clean.items():
        if cleaned.rstrip("s") == singular:
            return ct
    if not partial:
        return None
    candidates = [(cleaned, ct) for cleaned, ct in by_clean.items() if cleaned and (cleaned in target or target in cleaned)]
    return max(candidates, key=lambda p: len(p[0]))[1] if candidates else None


def resolve_sheet_check_types(sheet_names: List[str], check_types: List[str]) -> List[Tuple[str, str, str]]:
    """Return (sheet_name, check_type, kind). Exact matches win; partial matches only claim
    check types no exact sheet already claimed (so 'Tractor Hire' never overwrites 'Tractor')."""
    parsed = [(name, *split_sheet_name(name)) for name in sheet_names]
    resolved = {}
    for name, base, kind in parsed:
        ct = match_check_type(base, check_types)
        if ct:
            resolved[name] = ct
    claimed = {(ct, kind) for name, base, kind in parsed if name in resolved for ct in [resolved[name]]}
    for name, base, kind in parsed:
        if name in resolved:
            continue
        ct = match_check_type(base, check_types, partial=True)
        resolved[name] = ct if ct and (ct, kind) not in claimed else base
    return [(name, resolved[name], kind) for name, base, kind in parsed]


def _is_header_cell(value) -> bool:
    if value is None:
        return False
    h = clean_name(str(value))
    return h in ITEM_HEADERS or h.startswith("compuls") or h.startswith("compol") or h in ("critical", "common", "photo", "photorequired")


def _flag(row, col) -> bool:
    if col is None or len(row) <= col or row[col] is None:
        return False
    return str(row[col]).strip().lower() in ("yes", "y", "true", "1", "x", "compulsory")


def _row_text(row, item_col) -> str:
    if item_col is not None:
        cell = row[item_col] if len(row) > item_col else None
        return str(cell).strip() if cell is not None else ""
    for cell in row:
        if isinstance(cell, str) and cell.strip() and not re.fullmatch(r"[\d.\s]+", cell):
            return cell.strip()
    return ""


def parse_items(sheet) -> List[Dict]:
    """Read one template tab. Header row is optional; without it the first text column is the item."""
    rows = [r for r in sheet.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = rows[0]
    has_header = any(_is_header_cell(c) for c in header)
    item_col = compulsory_col = critical_col = photo_col = None
    if has_header:
        for i, cell in enumerate(header):
            h = clean_name(str(cell)) if cell is not None else ""
            if not h:
                continue
            if h in ITEM_HEADERS and item_col is None:
                item_col = i
            elif h.startswith("compuls") or h.startswith("compol"):
                compulsory_col = i
            elif h in ("critical", "common"):
                critical_col = i
            elif h.startswith("photo"):
                photo_col = i
        if item_col is None:
            item_col = 0
        rows = rows[1:]
    items = []
    for row in rows:
        text = _row_text(row, item_col)
        if len(text) < 2 or text.lower() in JUNK_ITEMS or re.fullmatch(r"[\d.\s]+", text):
            continue
        items.append({
            "item": text,
            "compulsory": _flag(row, compulsory_col),
            "critical": _flag(row, critical_col),
            "photo_required": _flag(row, photo_col),
        })
    return items


def parse_assets_sheet(sheet) -> List[Dict]:
    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]
    check_type_col = name_col = make_col = None
    for i, header in enumerate(headers):
        if header == "check type" or "checktype" in header:
            check_type_col = i
        elif header == "namecolumn" or ("name" in header and "check" not in header):
            name_col = i
        elif header == "makecolumn" or "make" in header:
            make_col = i
    if check_type_col is None or name_col is None or make_col is None:
        raise ValueError(f"Could not find Check Type, Name and Make columns in the file (found: {headers})")
    assets = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) <= max(check_type_col, name_col, make_col):
            continue
        check_type = str(row[check_type_col]).strip() if row[check_type_col] else ""
        name = str(row[name_col]).strip() if row[name_col] else ""
        make = str(row[make_col]).strip() if row[make_col] else ""
        if check_type and name and make:
            assets.append({"check_type": check_type, "name": name, "make": make})
    return assets


def parse_asset_workbook(file_content: bytes) -> Dict:
    workbook = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
    assets = parse_assets_sheet(workbook[workbook.sheetnames[0]])
    check_types = sorted({a["check_type"] for a in assets})
    checklist_templates: List[Dict] = []
    service_templates: List[Dict] = []
    processed_sheets: List[str] = []
    for sheet_name, check_type, kind in resolve_sheet_check_types(workbook.sheetnames[1:], check_types):
        items = parse_items(workbook[sheet_name])
        if not items:
            processed_sheets.append(f"{sheet_name} -> skipped (no items found)")
            continue
        if kind == "service":
            service_templates.append({
                "id": str(uuid.uuid4()),
                "check_type": check_type,
                "name": f"{check_type} Pre Service Check",
                "sheet_name": sheet_name,
                "sections": [i["item"] for i in items],
            })
            processed_sheets.append(f"{sheet_name} -> {check_type} (Pre Service Sheet: {len(items)} sections)")
        else:
            compulsory = sum(1 for i in items if i["compulsory"])
            checklist_templates.append({
                "id": str(uuid.uuid4()),
                "check_type": check_type,
                "sheet_name": sheet_name,
                "items": items,
            })
            processed_sheets.append(f"{sheet_name} -> {check_type} (Daily Check: {len(items)} items, {compulsory} compulsory)")
    return {
        "assets": assets,
        "checklist_templates": checklist_templates,
        "service_templates": service_templates,
        "processed_sheets": processed_sheets,
    }
